# -*- coding: utf-8 -*-
"""HTTP-level tests for the API layer (routes.py), via FastAPI's TestClient.

Complements the lower-level unit tests (test_excel_io.py etc.) by verifying
actual status codes, response shapes, and the reliability/observability
additions from V1.05 (file size limit, debug mode, friendly error bodies).
"""

import base64
from io import BytesIO

import openpyxl
import pytest
from fastapi.testclient import TestClient

from app.main import app

client = TestClient(app)


def _xlsx_file_tuple(xlsx_bytes: bytes, filename: str = "test.xlsx"):
    return ("file", (filename, BytesIO(xlsx_bytes), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"))


def test_health():
    res = client.get("/api/health")
    assert res.status_code == 200
    assert res.json() == {"status": "ok"}


def test_version():
    res = client.get("/api/version")
    assert res.status_code == 200
    assert "version" in res.json()


def test_convert_success_no_file_base64(sample_xlsx_bytes):
    res = client.post("/api/convert", files=[_xlsx_file_tuple(sample_xlsx_bytes)])
    assert res.status_code == 200
    data = res.json()
    assert "file_base64" not in data
    assert data["debug"] is None
    assert data["total_rows"] == len(data["rows"])
    assert data["summary"]["ppid_count"] == 2


def test_convert_debug_mode(sample_xlsx_bytes):
    res = client.post("/api/convert?debug=true", files=[_xlsx_file_tuple(sample_xlsx_bytes)])
    assert res.status_code == 200
    debug = res.json()["debug"]
    assert debug is not None
    assert "read_and_parse" in debug["stages_seconds"]
    assert "transform" in debug["stages_seconds"]
    assert debug["engine_used"] in ("calamine", "openpyxl (calamine fallback)")
    assert debug["peak_memory_mb"] > 0


def test_convert_rejects_wrong_extension(sample_xlsx_bytes):
    res = client.post("/api/convert", files=[("file", ("test.txt", BytesIO(sample_xlsx_bytes), "text/plain"))])
    assert res.status_code == 400
    assert "detail" in res.json()


def test_convert_rejects_oversized_upload():
    from app.api.routes import MAX_UPLOAD_SIZE_BYTES

    huge = b"PK\x03\x04" + b"0" * (MAX_UPLOAD_SIZE_BYTES + 1)
    res = client.post("/api/convert", files=[_xlsx_file_tuple(huge, "huge.xlsx")])
    assert res.status_code == 413
    assert "too large" in res.json()["detail"].lower()


def test_convert_rejects_empty_file():
    res = client.post("/api/convert", files=[_xlsx_file_tuple(b"", "empty.xlsx")])
    assert res.status_code == 400


def test_convert_rejects_garbage_content():
    res = client.post("/api/convert", files=[_xlsx_file_tuple(b"not an excel file", "fake.xlsx")])
    assert res.status_code == 400
    assert "detail" in res.json()


def test_convert_text_success():
    text = "PPID\tParameter\tREF.xxx\nX1\tPPID\tX1\nX1\tTS#1_CardName\tCARD_X\n"
    res = client.post("/api/convert-text", json={"text": text})
    assert res.status_code == 200
    assert res.json()["total_rows"] == 1


def test_convert_text_rejects_empty():
    res = client.post("/api/convert-text", json={"text": ""})
    assert res.status_code == 400


def test_export_success():
    rows = [{"PPID": "X1", "TS#": "TS#1", "CardName": "CARD_X", "FilmMaterial": "-",
             "CorrelationCard_1": "-", "CorrelationCard_2": "-", "CorrelationCard_3": "-",
             "DataCombination": "-", "DataFeedFoward": "-"}]
    res = client.post("/api/export", json={"filename": "out.xlsx", "rows": rows, "rule": None})
    assert res.status_code == 200
    data = res.json()
    assert data["filename"].endswith(".xlsx")
    xlsx_bytes = base64.b64decode(data["file_base64"])
    wb = openpyxl.load_workbook(BytesIO(xlsx_bytes))
    ws = wb.active
    header = next(ws.iter_rows(values_only=True))
    assert header[0] == "PPID"


def test_export_rejects_no_rows():
    res = client.post("/api/export", json={"filename": "out.xlsx", "rows": [], "rule": None})
    assert res.status_code == 400


def test_export_handles_malformed_rows_gracefully():
    """Rows missing expected keys shouldn't 500 - pandas fills them as NaN,
    which the export path must still turn into a valid file."""
    rows = [{"PPID": "X1"}]  # missing every other expected column
    res = client.post("/api/export", json={"filename": "out.xlsx", "rows": rows, "rule": None})
    assert res.status_code == 200


@pytest.mark.parametrize("rule_payload", [
    {"output_columns": ["PPID", "NotAField"], "column_order": ["NotAField", "PPID"]},
    {"output_columns": [], "column_order": []},
])
def test_export_handles_odd_rules_gracefully(rule_payload):
    rows = [{"PPID": "X1", "TS#": "TS#1", "CardName": "C", "FilmMaterial": "-",
             "CorrelationCard_1": "-", "CorrelationCard_2": "-", "CorrelationCard_3": "-",
             "DataCombination": "-", "DataFeedFoward": "-"}]
    res = client.post("/api/export", json={"filename": "out.xlsx", "rows": rows, "rule": rule_payload})
    assert res.status_code == 200


def _description_xlsx_tuple(rows: list[list], filename: str = "desc.xlsx"):
    buf = BytesIO()
    wb = openpyxl.Workbook()
    ws = wb.active
    for row in rows:
        ws.append(row)
    wb.save(buf)
    return ("file", (filename, BytesIO(buf.getvalue()), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"))


_CONVERTED_ROWS = [
    {"PPID": "AB000010_1", "TS#": "TS#1", "CardName": "CARD1", "FilmMaterial": "-",
     "CorrelationCard_1": "-", "CorrelationCard_2": "-", "CorrelationCard_3": "-",
     "DataCombination": "-", "DataFeedFoward": "-"},
    {"PPID": "AB000020_1", "TS#": "TS#1", "CardName": "CARDX", "FilmMaterial": "-",
     "CorrelationCard_1": "-", "CorrelationCard_2": "-", "CorrelationCard_3": "-",
     "DataCombination": "-", "DataFeedFoward": "-"},
]


def test_export_includes_desc_when_present():
    rows = [{**_CONVERTED_ROWS[0], "DESC": "First PPID"}]
    res = client.post("/api/export", json={"filename": "out.xlsx", "rows": rows, "rule": None})
    assert res.status_code == 200
    xlsx_bytes = base64.b64decode(res.json()["file_base64"])
    wb = openpyxl.load_workbook(BytesIO(xlsx_bytes))
    ws = wb.active
    header = next(ws.iter_rows(values_only=True))
    assert "DESC" in header


def test_export_places_desc_immediately_after_ppid():
    rows = [{**_CONVERTED_ROWS[0], "DESC": "First PPID"}]
    res = client.post("/api/export", json={"filename": "out.xlsx", "rows": rows, "rule": None})
    assert res.status_code == 200
    xlsx_bytes = base64.b64decode(res.json()["file_base64"])
    wb = openpyxl.load_workbook(BytesIO(xlsx_bytes))
    header = next(wb.active.iter_rows(values_only=True))
    assert header[0] == "PPID"
    assert header[1] == "DESC"


def test_export_places_desc_after_ppid_even_with_a_reordering_rule():
    rows = [{**_CONVERTED_ROWS[0], "DESC": "First PPID"}]
    rule = {"output_columns": ["PPID", "CardName", "TS#"], "column_order": ["CardName", "PPID", "TS#"]}
    res = client.post("/api/export", json={"filename": "out.xlsx", "rows": rows, "rule": rule})
    assert res.status_code == 200
    xlsx_bytes = base64.b64decode(res.json()["file_base64"])
    header = list(next(openpyxl.load_workbook(BytesIO(xlsx_bytes)).active.iter_rows(values_only=True)))
    assert header == ["CardName", "PPID", "DESC", "TS#"]


def test_add_description_success():
    import json

    desc_file = _description_xlsx_tuple([["PPID", "DESC"], ["AB000010_1", "First PPID"], ["AB000020_1", "Second PPID"]])
    res = client.post(
        "/api/add-description",
        files=[desc_file],
        data={"rows": json.dumps(_CONVERTED_ROWS)},
    )
    assert res.status_code == 200
    data = res.json()
    assert data["matched_count"] == 2
    assert data["unmatched_count"] == 0
    assert all(row["DESC"] for row in data["rows"])
    assert data["columns"][:2] == ["PPID", "DESC"]
    assert list(data["rows"][0].keys())[:2] == ["PPID", "DESC"]


def test_add_description_reports_unmatched():
    import json

    desc_file = _description_xlsx_tuple([["PPID", "DESC"], ["AB000010_1", "First PPID"]])
    res = client.post(
        "/api/add-description",
        files=[desc_file],
        data={"rows": json.dumps(_CONVERTED_ROWS)},
    )
    assert res.status_code == 200
    data = res.json()
    assert data["matched_count"] == 1
    assert data["unmatched_count"] == 1
    assert data["unmatched_ppids"] == ["AB000020_1"]


def test_add_description_rejects_duplicate_ppid():
    import json

    desc_file = _description_xlsx_tuple([["PPID", "DESC"], ["AB000010_1", "A"], ["AB000010_1", "B"]])
    res = client.post(
        "/api/add-description",
        files=[desc_file],
        data={"rows": json.dumps(_CONVERTED_ROWS)},
    )
    assert res.status_code == 400
    assert "duplicate" in res.json()["detail"].lower()


def test_add_description_rejects_missing_desc_column():
    import json

    desc_file = _description_xlsx_tuple([["PPID", "Notes"], ["AB000010_1", "irrelevant"]])
    res = client.post(
        "/api/add-description",
        files=[desc_file],
        data={"rows": json.dumps(_CONVERTED_ROWS)},
    )
    assert res.status_code == 400


def test_add_description_rejects_empty_rows():
    import json

    desc_file = _description_xlsx_tuple([["PPID", "DESC"], ["AB000010_1", "A"]])
    res = client.post(
        "/api/add-description",
        files=[desc_file],
        data={"rows": json.dumps([])},
    )
    assert res.status_code == 400
